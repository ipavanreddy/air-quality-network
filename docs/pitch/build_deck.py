"""Build the VayuDrishti pitch deck (12 slides, 16:9).

    uv run --with python-pptx python docs/pitch/build_deck.py

Writes docs/pitch/VayuDrishti_pitch.pptx. Content is structured around the hackathon evaluation criteria
(AI/technical 25 %, problem-solution fit 20 %, depth & reach 20 %, deployability & scalability 20 %,
impact 15 %) and ends with cross-border (BRICS) portability. Live URLs for the apps are filled in by the lead.
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).with_name("VayuDrishti_pitch.pptx")

INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x65, 0x72)
TEAL = RGBColor(0x0F, 0x4C, 0x5C)
TEAL_LIGHT = RGBColor(0xE3, 0xEF, 0xF1)
AMBER = RGBColor(0xE3, 0x64, 0x14)
AMBER_LIGHT = RGBColor(0xFD, 0xEB, 0xDD)
GREEN = RGBColor(0x2D, 0x6A, 0x4F)
GREEN_LIGHT = RGBColor(0xE2, 0xF0, 0xE8)
GREY_LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Calibri"

API_URL = "https://air-quality-network-api-847963771142.asia-south1.run.app"
REPO_URL = "https://github.com/ipavanreddy/air-quality-network"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W = prs.slide_width


def text(slide, x, y, w, h, content, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, font=FONT):
    """Add a text box. `content` is a string or a list of strings / (string, overrides) paragraphs."""
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = content if isinstance(content, list) else [content]
    for i, p in enumerate(paras):
        opts = {}
        if isinstance(p, tuple):
            p, opts = p
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = opts.get("align", align)
        para.space_after = Pt(opts.get("space_after", 6))
        run = para.add_run()
        run.text = p
        run.font.name = font
        run.font.size = Pt(opts.get("size", size))
        run.font.bold = opts.get("bold", bold)
        run.font.color.rgb = opts.get("color", color)
    return box


def bullets(slide, x, y, w, h, items, size=16, color=INK, bullet="•"):
    return text(slide, x, y, w, h, [f"{bullet}  {i}" for i in items], size=size, color=color)


def card(slide, x, y, w, h, title, body, fill=TEAL_LIGHT, accent=TEAL, title_size=16, body_size=13):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.08
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y + Inches(0.18), Inches(0.07), h - Inches(0.36))
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    text(slide, x + Inches(0.2), y + Inches(0.12), w - Inches(0.3), Inches(0.5), title, size=title_size,
         bold=True, color=accent)
    body_items = body if isinstance(body, list) else [body]
    text(slide, x + Inches(0.2), y + Inches(0.6), w - Inches(0.3), h - Inches(0.7), body_items,
         size=body_size, color=INK)
    return shape


def pill(slide, x, y, w, h, label, fill=TEAL, color=WHITE, size=13, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shape.adjustments[0] = 0.3
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = line
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return shape


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd",
                          {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


def header(slide, title, kicker=None, criterion=None):
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    band.fill.solid()
    band.fill.fore_color.rgb = TEAL
    band.line.fill.background()
    if kicker:
        text(slide, Inches(0.6), Inches(0.35), Inches(9), Inches(0.4), kicker.upper(), size=12, bold=True,
             color=AMBER)
    text(slide, Inches(0.6), Inches(0.62), Inches(10.5), Inches(0.9), title, size=30, bold=True, color=INK)
    if criterion:
        pill(slide, Inches(10.6), Inches(0.42), Inches(2.2), Inches(0.42), criterion, fill=AMBER_LIGHT,
             color=AMBER, size=12)


def footer(slide, n):
    text(slide, Inches(0.6), Inches(7.0), Inches(8), Inches(0.35),
         "VayuDrishti · Build with AI (Google) · Track 2", size=10, color=MUTED)
    text(slide, Inches(11.9), Inches(7.0), Inches(0.9), Inches(0.35), str(n), size=10, color=MUTED,
         align=PP_ALIGN.RIGHT)


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


slides = []


def new_slide():
    s = prs.slides.add_slide(BLANK)
    slides.append(s)
    return s


# 1 · Title ------------------------------------------------------------------------------------------
s = new_slide()
bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, prs.slide_height)
bg.fill.solid()
bg.fill.fore_color.rgb = TEAL
bg.line.fill.background()
text(s, Inches(0.9), Inches(1.5), Inches(11), Inches(0.5), "BUILD WITH AI (GOOGLE) · TRACK 2 · CLIMATE ACTION",
     size=14, bold=True, color=AMBER_LIGHT)
text(s, Inches(0.9), Inches(2.1), Inches(11), Inches(1.3), "VayuDrishti", size=66, bold=True, color=WHITE)
text(s, Inches(0.9), Inches(3.35), Inches(11.5), Inches(1.0),
     "From a citizen's photo to an officer's action, in seconds.", size=30, color=WHITE)
text(s, Inches(0.9), Inches(4.4), Inches(11.2), Inches(1.2),
     "Federated, hyper-local air-quality intelligence. Gemini fuses citizen photos, low-cost sensors, "
     "satellite and weather data into evidence-backed alerts for the right authority, in English, "
     "Hindi and Punjabi.", size=18, color=TEAL_LIGHT)
text(s, Inches(0.9), Inches(6.2), Inches(11.5), Inches(0.8),
     [f"Code: {REPO_URL}", f"API: {API_URL}   ·   Apps: Vercel (links in README)"], size=13, color=TEAL_LIGHT)
notes(s, "One sentence: VayuDrishti finds pollution before the monitors do and tells the officer who can stop it.")

# 2 · Problem ----------------------------------------------------------------------------------------
s = new_slide()
header(s, "Cities see the average, not the 3 a.m. smoke next door", "The problem", "Problem fit · 20%")
card(s, Inches(0.6), Inches(1.8), Inches(6.0), Inches(2.3), "What officers have today", [
    "A sparse network of official stations and city-wide averages",
    "Manual complaints that are hard to verify",
    "Enforcement after pollution has peaked",
    "Little coordination when smoke crosses state lines",
], fill=GREY_LIGHT, accent=MUTED, body_size=14)
card(s, Inches(6.9), Inches(1.8), Inches(5.8), Inches(2.3), "What already exists, but disconnected", [
    "Citizens with phones who see smoke, burning and dust",
    "Low-cost community sensors (uncalibrated)",
    "Satellites: fires, NO2, aerosols. Weather: wind, humidity",
], fill=AMBER_LIGHT, accent=AMBER, body_size=14)
pill(s, Inches(0.6), Inches(4.45), Inches(12.1), Inches(1.1),
     "“The station 6 km away reads ‘Poor’. Beside the industrial estate, residents breathe far "
     "worse air at 3 a.m., and nobody is alerted.”", fill=TEAL, size=18, bold=False)
bullets(s, Inches(0.6), Inches(5.8), Inches(12), Inches(1.1), [
    "Result: undetected sources, badly targeted inspections, no early warning, direct harm to public health.",
], size=15, color=MUTED)
footer(s, 2)

# 3 · Solution: one journey -----------------------------------------------------------------------------
s = new_slide()
header(s, "One polished journey: evidence in, approved action out", "The solution", "Problem fit · 20%")
steps = [
    ("Citizen photo\n+ voice note", AMBER), ("Gemini\nverification", TEAL), ("Sensors + satellite\n+ weather on grid", TEAL),
    ("Hotspot\nConfidence", TEAL), ("Likely source\n+ action brief", TEAL), ("Officer alert\n(human approval)", GREEN),
]
x0, y0, bw, bh, gap = Inches(0.6), Inches(2.0), Inches(1.85), Inches(1.2), Inches(0.2)
for i, (label, col) in enumerate(steps):
    x = x0 + i * (bw + gap)
    pill(s, x, y0, bw, bh, label, fill=col, size=14)
    if i < len(steps) - 1:
        arrow(s, x + bw, y0 + bh // 2, x + bw + gap, y0 + bh // 2)
pill(s, Inches(3.4), Inches(3.55), Inches(3.0), Inches(0.8), "72 h corridor forecast", fill=TEAL_LIGHT, color=TEAL)
pill(s, Inches(6.7), Inches(3.55), Inches(3.3), Inches(0.8), "Voice advisory: en · हिन्दी · ਪੰਜਾਬੀ", fill=TEAL_LIGHT,
     color=TEAL)
pill(s, Inches(10.3), Inches(3.55), Inches(2.4), Inches(0.8), "Cross-boundary alert", fill=GREEN_LIGHT, color=GREEN)
card(s, Inches(0.6), Inches(4.75), Inches(3.9), Inches(1.95), "Citizen Reporter", [
    "Photo, voice or typed report with a Maps-confirmed location",
    "Instant AI verdict and status",
    "Local forecast and spoken advisories",
], body_size=13)
card(s, Inches(4.7), Inches(4.75), Inches(3.9), Inches(1.95), "Environmental Officer", [
    "1 km hotspot map and alert inbox",
    "Evidence, likely source, recommended inspection",
    "Approve → record action → close",
], body_size=13)
card(s, Inches(8.8), Inches(4.75), Inches(3.9), Inches(1.95), "Neighbouring states", [
    "Same canonical schema",
    "Shared, fine-tuned models",
    "Officer-approved cross-boundary alerts",
], fill=GREEN_LIGHT, accent=GREEN, body_size=13)
footer(s, 3)

# 4 · Demo storyline ------------------------------------------------------------------------------------
s = new_slide()
header(s, "What the live demo shows (5 minutes)", "Demo", None)
rows = [
    ("0:30", "Citizen report", "Hindi voice note → Speech-to-Text; Gemini: industrial emission, severity 4/5, ~90 %"),
    ("1:10", "Evidence fusion", "Hotspot Confidence 66 → ~84 (threshold 70): sensor 35 · satellite 25 · citizen 25 · weather 15"),
    ("1:50", "Action brief", "Likely source + evidence with source/time + recommended inspection + uncertainties"),
    ("2:30", "Forecast", "Now / +24 / +48 / +72 h slider, range and drivers, model vs persistence"),
    ("3:00", "Officer action", "Alert in ≤ 5 s → acknowledge & approve → record action → close; audit trail"),
    ("3:30", "Advisory", "Gemini draft → Cloud Translation → Text-to-Speech in Punjabi, after approval"),
    ("3:50", "Interoperability", "Punjab stubble fire → cross-boundary alert to Delhi; shared model card"),
]
y = Inches(1.75)
for t, name, desc in rows:
    pill(s, Inches(0.6), y, Inches(0.9), Inches(0.55), t, fill=AMBER, size=13)
    text(s, Inches(1.7), y + Inches(0.05), Inches(2.4), Inches(0.5), name, size=16, bold=True, color=TEAL)
    text(s, Inches(4.1), y + Inches(0.07), Inches(8.7), Inches(0.5), desc, size=14)
    y += Inches(0.72)
footer(s, 4)
notes(s, "Numbers come from a live run on Vertex AI; they move by a point or two between runs.")

# 5 · AI & technical execution ---------------------------------------------------------------------------
s = new_slide()
header(s, "Gemini does real work, and is kept honest", "AI & technical execution", "AI / tech · 25%")
card(s, Inches(0.6), Inches(1.8), Inches(3.9), Inches(2.45), "Multimodal verification", [
    "Gemini 2.5 Flash on Vertex AI reads the photo",
    "Event yes/no, source type, severity, indicators, confidence",
    "Irrelevant or poor images are flagged",
], body_size=15)
card(s, Inches(4.7), Inches(1.8), Inches(3.9), Inches(2.45), "Grounded action briefs", [
    "Input: only the structured evidence bundle",
    "Actions limited to each jurisdiction's rules",
    "Any number not in the data is flagged",
], body_size=15)
card(s, Inches(8.8), Inches(1.8), Inches(3.9), Inches(2.45), "Language and voice", [
    "Gemini advisory → Cloud Translation (hi, pa)",
    "Text-to-Speech playback, Speech-to-Text notes",
    "Maps Geocoding confirms the location",
], body_size=15)
card(s, Inches(0.6), Inches(4.45), Inches(6.0), Inches(2.3), "Transparent scoring, not a black box", [
    "Hotspot Confidence = sensor anomaly 35 + satellite 25 + citizen 25 + weather plausibility 15",
    "κ-Köhler humidity calibration on low-cost sensors (kept as indicative)",
    "Per-horizon forecast model, backtested against persistence",
], fill=AMBER_LIGHT, accent=AMBER, body_size=15)
card(s, Inches(6.9), Inches(4.45), Inches(5.8), Inches(2.3), "Structured AI orchestration", [
    "Versioned prompts + JSON schemas (ai/prompts, ai/schemas)",
    "model_name, model_version, prompt_version on every AI record",
    "Retries, low temperature, labelled fallback when AI is unavailable",
], fill=GREEN_LIGHT, accent=GREEN, body_size=15)
footer(s, 5)

# 6 · Architecture -----------------------------------------------------------------------------------------
s = new_slide()
header(s, "Cloud-native and API-first on Google Cloud", "Architecture", "Deployability · 20%")
pill(s, Inches(0.6), Inches(1.9), Inches(2.9), Inches(1.0), "Citizen app\nNext.js · Vercel", fill=AMBER, size=14)
pill(s, Inches(0.6), Inches(3.3), Inches(2.9), Inches(1.0), "Officer dashboard\nNext.js · Vercel", fill=AMBER, size=14)
core = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.2), Inches(1.7), Inches(4.3), Inches(4.9))
core.adjustments[0] = 0.05
core.fill.solid()
core.fill.fore_color.rgb = TEAL_LIGHT
core.line.color.rgb = TEAL
text(s, Inches(4.4), Inches(1.8), Inches(4), Inches(0.5), "Cloud Run · FastAPI (asia-south1)", size=15, bold=True,
     color=TEAL)
mods = ["reports + AI verification", "sensors + calibration", "geospatial (Earth Engine → H3)", "hotspots + confidence",
        "forecasting + model registry", "alerts + cross-boundary", "localization (voice, translation)",
        "interop: adapters + canonical schema"]
for i, m in enumerate(mods):
    pill(s, Inches(4.45), Inches(2.35) + i * Inches(0.5), Inches(3.8), Inches(0.4), m, fill=WHITE, color=INK, size=12,
         bold=False)
svc = [("Gemini 2.5 Flash · Vertex AI", TEAL), ("Translation · TTS · STT", TEAL), ("Maps Geocoding + JS", TEAL),
       ("BigQuery · Cloud Storage", GREEN), ("Earth Engine (wired)", MUTED), ("Secret Manager · Cloud Build", MUTED)]
for i, (label, col) in enumerate(svc):
    pill(s, Inches(9.3), Inches(1.8) + i * Inches(0.8), Inches(3.4), Inches(0.62), label, fill=col, size=13)
arrow(s, Inches(3.5), Inches(2.4), Inches(4.2), Inches(2.9))
arrow(s, Inches(3.5), Inches(3.8), Inches(4.2), Inches(3.8))
arrow(s, Inches(8.5), Inches(4.1), Inches(9.3), Inches(4.1))
text(s, Inches(0.6), Inches(4.7), Inches(3.4), Inches(1.9), [
    "Adapters (JSON):", "Delhi NCR · Punjab · Maharashtra", "→ canonical schema",
], size=13, color=MUTED)
footer(s, 6)

# 7 · Trust ---------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Built for officers who must defend every action", "Trust & safety", "AI / tech · 25%")
card(s, Inches(0.6), Inches(1.8), Inches(3.9), Inches(2.3), "Human approval", [
    "Alerts start as pending review",
    "Officer approves actions, transfers and advisories",
    "Full audit trail",
], body_size=14)
card(s, Inches(4.7), Inches(1.8), Inches(3.9), Inches(2.3), "Evidence first", [
    "Observed data → AI interpretation → recommendation",
    "Source + timestamp on every signal",
    "Data freshness and uncertainties shown",
], body_size=14)
card(s, Inches(8.8), Inches(1.8), Inches(3.9), Inches(2.3), "Careful language", [
    "“Likely source” and “recommended inspection”",
    "No named companies, farms or people",
    "Low-cost sensors labelled indicative",
], body_size=14)
card(s, Inches(0.6), Inches(4.35), Inches(12.1), Inches(2.3), "Honest demo mode", [
    "Every integration reports real / demo / fallback at /api/config, shown as a badge in both apps.",
    "In production: live Gemini, Translation, TTS, STT, Maps, BigQuery and Cloud Storage. Earth Engine, CPCB and "
    "sensor feeds run on clearly labelled sample data until credentials arrive.",
    "Pseudonymous reporters; the photo store is separate from public views.",
], fill=AMBER_LIGHT, accent=AMBER, body_size=14)
footer(s, 7)

# 8 · Depth & reach ------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Three states, three problems, one platform", "Depth & reach across India", "Depth & reach · 20%")
states = [
    ("Delhi NCR", "Industrial + traffic smog", ["Threshold 70", "Shared IGP smog model", "Hindi + English"], TEAL,
     TEAL_LIGHT),
    ("Punjab", "Crop-residue burning", ["~40 upwind fires (FIRMS)", "Cross-boundary → Delhi", "Punjabi + Hindi"], AMBER,
     AMBER_LIGHT),
    ("Maharashtra", "Construction dust", ["Threshold 65", "Persistence baseline model", "Different feed format (UTC)"],
     GREEN, GREEN_LIGHT),
]
for i, (name, problem, pts, acc, fill) in enumerate(states):
    card(s, Inches(0.6) + i * Inches(4.1), Inches(1.8), Inches(3.9), Inches(2.6), name, [problem, *pts], fill=fill,
         accent=acc, body_size=14)
bullets(s, Inches(0.6), Inches(4.7), Inches(12.1), Inches(2.2), [
    "Satellite and weather layers (Sentinel-5P, FIRMS, MODIS MAIAC, ERA5-Land, GFS) cover every Indian district "
    "from day one, before any local sensor is connected.",
    "Three languages today (English, हिन्दी, ਪੰਜਾਬੀ). Adding a language is a translation target and a voice locale.",
    "H3 grid at ~1 km: the same cell ids work for any city, so evidence and models travel.",
], size=15)
footer(s, 8)

# 9 · Interoperability ---------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Federated by design: share models, not raw data", "Interoperability", "Depth & reach · 20%")
pill(s, Inches(0.6), Inches(1.9), Inches(3.3), Inches(0.9), "Punjab feed\ndevice · time_ist · PM2_5 · HUM", fill=AMBER,
     size=12)
pill(s, Inches(0.6), Inches(3.0), Inches(3.3), Inches(0.9), "Delhi feed\nsensor_id · ts · pm25 · rh", fill=TEAL, size=12)
pill(s, Inches(0.6), Inches(4.1), Inches(3.3), Inches(0.9), "Maharashtra feed\nid · timestamp_utc · pm2_5_ugm3",
     fill=GREEN, size=12)
pill(s, Inches(4.6), Inches(3.0), Inches(2.6), Inches(0.9), "State adapters\n(field map + time zone)", fill=GREY_LIGHT,
     color=INK, size=13)
pill(s, Inches(7.9), Inches(3.0), Inches(2.3), Inches(0.9), "Canonical cell\nsnapshot schema", fill=TEAL, size=13)
for yy in (2.35, 3.45, 4.55):
    arrow(s, Inches(3.9), Inches(yy), Inches(4.6), Inches(3.45))
arrow(s, Inches(7.2), Inches(3.45), Inches(7.9), Inches(3.45))
card(s, Inches(10.4), Inches(1.9), Inches(2.3), Inches(3.1), "Shared services", [
    "Hotspot scoring", "AI verification + briefs", "Forecast model", "Alert exchange",
], body_size=13)
card(s, Inches(0.6), Inches(5.3), Inches(6.0), Inches(1.5), "Shared model with a model card", [
    "igp-smog-forecast-v1: pooled on Delhi NCR + Punjab, fine-tuned per state, MAE vs persistence published",
], fill=AMBER_LIGHT, accent=AMBER, body_size=13)
card(s, Inches(6.9), Inches(5.3), Inches(5.8), Inches(1.5), "Cross-boundary alert", [
    "Punjab officer approves → Delhi NCR inbox receives it with ETA and evidence",
], fill=GREEN_LIGHT, accent=GREEN, body_size=13)
footer(s, 9)

# 10 · Deployability & scalability ------------------------------------------------------------------------------------
s = new_slide()
header(s, "Deployable today, scalable by configuration", "Deployability & scalability", "Deployability · 20%")
card(s, Inches(0.6), Inches(1.8), Inches(5.9), Inches(2.5), "Running now", [
    "API on Cloud Run (asia-south1) with a least-privilege service account",
    "Two Next.js apps on Vercel, each builds on its own",
    "Keys in Secret Manager; one-command deploy script",
    "BigQuery analytics mirror, Cloud Storage evidence",
], body_size=14)
card(s, Inches(6.8), Inches(1.8), Inches(5.9), Inches(2.5), "Onboarding a new state = one JSON file", [
    "Jurisdiction, threshold, languages, action rules",
    "Sensor field map → canonical schema",
    "Pick or fine-tune a shared model",
    "Declare downwind neighbours",
], fill=GREEN_LIGHT, accent=GREEN, body_size=14)
phases = [("Hackathon", "3 states · 1 corridor · 3 languages"), ("City pilot", "10–20 co-located sensors · field teams"),
          ("State", "many cities · state models · more languages"),
          ("National", "shared registry · federated learning")]
for i, (p, d) in enumerate(phases):
    x = Inches(0.6) + i * Inches(3.1)
    pill(s, x, Inches(4.7), Inches(2.8), Inches(0.6), p, fill=TEAL if i == 0 else GREY_LIGHT,
         color=WHITE if i == 0 else TEAL, size=14)
    text(s, x, Inches(5.4), Inches(2.8), Inches(0.9), d, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    if i < 3:
        arrow(s, x + Inches(2.8), Inches(5.0), x + Inches(3.1), Inches(5.0))
text(s, Inches(0.6), Inches(6.3), Inches(12), Inches(0.5),
     "Next scaling step: move the operational store to Firestore (already wired) and run many API instances.",
     size=13, color=MUTED)
footer(s, 10)

# 11 · Impact + BRICS ------------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Impact where air is worst, and portable beyond India", "Impact & portability", "Impact · 15%")
card(s, Inches(0.6), Inches(1.8), Inches(6.0), Inches(2.9), "Impact in India", [
    "Targets urban and peri-urban people in the most polluted corridors",
    "Faster, better-targeted inspections from verified evidence",
    "24–72 h warning for sensitive groups, in their language and by voice",
    "Coordinated action on transboundary smoke",
], body_size=16)
card(s, Inches(6.9), Inches(1.8), Inches(5.8), Inches(2.9), "How we would measure it", [
    "Time from first report to officer action",
    "Share of alerts confirmed on inspection",
    "PM2.5 change after intervention (action effectiveness)",
    "Advisory reach by language",
], fill=GREY_LIGHT, accent=MUTED, body_size=16)
card(s, Inches(0.6), Inches(4.95), Inches(12.1), Inches(1.85), "Cross-border / BRICS portability", [
    "Nothing is India-specific in the core: H3 grid, global satellite and reanalysis layers, a canonical schema and "
    "per-jurisdiction adapters. A partner city adds an adapter, its AQI breakpoints, languages and action rules, "
    "and can reuse or fine-tune shared models. It can exchange transboundary alerts without sharing raw data.",
], fill=AMBER_LIGHT, accent=AMBER, body_size=16)
footer(s, 11)

# 12 · Status & ask ------------------------------------------------------------------------------------------------------------
s = new_slide()
header(s, "Status, next steps, and links", "Where we are", None)
card(s, Inches(0.6), Inches(1.8), Inches(4.0), Inches(3.3), "Live in production", [
    "Gemini 2.5 Flash (Vertex AI)", "Translation · Text-to-Speech · Speech-to-Text", "Maps Geocoding + JS",
    "BigQuery · Cloud Storage · Cloud Run",
], fill=GREEN_LIGHT, accent=GREEN, body_size=14)
card(s, Inches(4.8), Inches(1.8), Inches(4.0), Inches(3.3), "Labelled sample, adapter ready", [
    "Earth Engine layers (registration pending)", "CPCB station feed (API key)", "Community sensor feeds",
    "Firestore, Vertex AI Model Registry",
], fill=AMBER_LIGHT, accent=AMBER, body_size=14)
card(s, Inches(9.0), Inches(1.8), Inches(3.7), Inches(3.3), "Next 90 days", [
    "City pilot with a clean-air cell", "Co-locate sensors for calibration", "Officer login + SMS alerts",
    "Action-effectiveness analytics",
], body_size=14)
text(s, Inches(0.6), Inches(5.45), Inches(12.1), Inches(1.4), [
    (f"Code: {REPO_URL}", {"size": 15}),
    (f"API: {API_URL}", {"size": 15}),
    ("Citizen app and officer dashboard: Vercel (URLs in README)", {"size": 15}),
], color=TEAL)
footer(s, 12)

prs.save(OUT)
print(f"wrote {OUT} ({len(prs.slides)} slides)")
